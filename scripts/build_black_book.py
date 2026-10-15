"""Build the editable black-book DOCX from verified repository artifacts."""

from __future__ import annotations

import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "final_submission" / "Track_A_Black_Book.docx"


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd"); element.set(qn("w:fill"), fill); props.append(element)


def borders(table) -> None:
    props = table._tbl.tblPr
    element = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}"); edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6"); edge.set(qn("w:color"), "D9D9D9"); element.append(edge)
    props.append(element)


def page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.3) -> None:
    if not path.exists():
        return
    paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.style = doc.styles["Caption"]


def add_results_table(doc: Document) -> None:
    path = ROOT / "reports/results/comparison/method_comparison_summary.csv"
    rows = list(csv.DictReader(path.open(encoding="utf-8")))
    table = doc.add_table(rows=1, cols=5); table.autofit = False; borders(table)
    headers = ["Method", "Precision", "Recall", "F1 mean SD", "ROC AUC"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value; shade(cell, "1E3A5F")
        for run in cell.paragraphs[0].runs: run.font.color.rgb = RGBColor(255, 255, 255); run.font.bold = True
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        method = "ML gradient boosting" if row["method"] == "ML (HistGradientBoosting)" else row["method"]
        values = [method, f'{float(row["precision_mean"]):.3f}', f'{float(row["recall_mean"]):.3f}',
                  f'{float(row["f1_mean"]):.3f} ({float(row["f1_std"]):.3f})', f'{float(row["roc_auc_mean"]):.3f}']
        for cell, value in zip(cells, values):
            cell.text = value; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index % 2: shade(cell, "EEF4FA")
    widths = [1.55, 1.0, 1.0, 1.3, 1.0]
    for row in table.rows:
        for cell, width in zip(row.cells, widths): cell.width = Inches(width)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.95)
    section.left_margin = section.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[name]; style.font.name = "Arial"; style.font.size = Pt(size)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = RGBColor(0, 0, 0); style.font.bold = True
        style.paragraph_format.space_before = Pt(14); style.paragraph_format.space_after = Pt(7)
    title_props = doc.styles["Title"]._element.get_or_add_pPr()
    border = title_props.find(qn("w:pBdr"))
    if border is not None:
        title_props.remove(border)
    doc.styles["Caption"].font.name = "Times New Roman"; doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.italic = True; doc.styles["Caption"].font.color.rgb = RGBColor(70, 70, 70)
    page_number(section.footer.paragraphs[0])


def build() -> Path:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    source = (ROOT / "reports/black_book.md").read_text(encoding="utf-8").splitlines()
    doc = Document(); configure(doc)
    title = doc.add_paragraph(style="Title"); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Inferring Hidden Economic Connections from Financial Shock Propagation")
    subtitle = doc.add_paragraph("A Synthetic and Real Data Prototype")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER; subtitle.runs[0].font.size = Pt(16)
    doc.add_paragraph(); metadata = doc.add_paragraph(); metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.add_run("University Black Book\n\nStudent: [STUDENT NAME AND ID]\nGuide: [GUIDE NAME]\nDepartment and institution: [INSTITUTIONAL DETAILS]\nAcademic year: [ACADEMIC YEAR]")
    doc.add_page_break()
    in_code = False; paragraph_buffer: list[str] = []

    def flush() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            text = " ".join(part.strip() for part in paragraph_buffer).strip()
            text = re.sub(r"\*\*(.*?)\*\*", r"\1", text).replace("`", "")
            replacements = {r"\(n\)": "n", r"\(A\)": "A", r"\(A_{ij}=A_{ji}\)": "Aij = Aji",
                r"\(A_{ii}=0\)": "Aii = 0", r"\(A_{ij}\in\{0,1\}\)": "Aij in {0, 1}",
                r"\(\alpha\)": "alpha", r"\(\beta\)": "beta", r"\(D^{-1}A\)": "D^-1 A",
                r"\(s_t\)": "s(t)", r"\(\epsilon_t\)": "epsilon(t)",
                r"\(\alpha+\beta<1\)": "alpha + beta < 1", r"\(e=(t,i,m)\)": "e = (t, i, m)",
                r"\(j\)": "j", r"\(L\)": "L", r"\(z\)": "z", r"\(\lambda\)": "lambda",
                r"\(\tau_r\)": "tau-r", r"\(Q_{ij}=(Q_{i\rightarrow j}+Q_{j\rightarrow i})/2\)": "Qij = (Qi-to-j + Qj-to-i) / 2",
                r"\(|z|\geq2.5\)": "|z| >= 2.5"}
            for old, new in replacements.items(): text = text.replace(old, new)
            text = text.replace("*", "")
            if text: doc.add_paragraph(text)
            paragraph_buffer = []

    for line in source:
        if line.startswith("# Inferring") or line.startswith("## A Synthetic") or line.startswith("**Student") or line.startswith("**Guide") or line.startswith("**Department") or line.startswith("**Academic"):
            continue
        if line.startswith("```"):
            flush(); in_code = not in_code; continue
        if in_code:
            continue
        if line.startswith("\\[") or line.startswith("\\]"):
            flush(); continue
        if line.strip().startswith("x_{"):
            flush(); p = doc.add_paragraph("x(t+1) = alpha x(t) + beta D^-1 A x(t) + s(t+1) + epsilon(t+1)")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.italic = True; continue
        if line.strip().startswith("r_{"):
            flush(); p = doc.add_paragraph("r(e, i-to-j) = max over lags [exp(-lambda(lag-1)) times max(0, sign(m) z(t+lag, j))]")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.italic = True; continue
        if line.strip().startswith("Q_"):
            flush(); p = doc.add_paragraph("Qi-to-j = mean event response times the fraction of responses at or above tau-r")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.italic = True; continue
        match = re.match(r"^(#{2,4})\s+(.*)", line)
        if match:
            flush(); level = min(len(match.group(1)) - 1, 3)
            heading = re.sub(r"\*\*(.*?)\*\*", r"\1", match.group(2))
            if level == 1 and heading not in ("Abstract",):
                doc.add_page_break()
            else:
                doc.add_paragraph()
            if heading == "7.1 Baseline comparison":
                doc.add_page_break()
            doc.add_heading(heading, level=level)
            if heading == "4 Proposed System Design":
                add_figure(doc, ROOT / "reports/figures/system_architecture.png", "Figure 1 System architecture")
                add_figure(doc, ROOT / "reports/figures/data_flow.png", "Figure 2 Synthetic and real-data flow", 6.0)
            elif heading == "7 Experimental Setup and Results":
                add_figure(doc, ROOT / "reports/figures/true_network.png", "Figure 3 Synthetic hidden network")
                add_figure(doc, ROOT / "reports/figures/inferred_network.png", "Figure 4 Inferred network")
            elif heading == "7.1 Baseline comparison":
                add_results_table(doc)
                doc.add_page_break()
                add_figure(doc, ROOT / "reports/results/comparison/method_comparison_f1.png", "Figure 5 Method comparison on independent test seeds", 5.8)
            elif heading == "7.2 Robustness":
                for number, stem, label in ((6, "noise_std", "noise"), (7, "propagation", "propagation strength"), (8, "shocks", "shock count")):
                    add_figure(doc, ROOT / f"reports/results/robustness/f1_vs_{stem}.png", f"Figure {number} F1 versus {label}", 5.2)
            elif heading == "7.3 Real-data demonstration":
                add_figure(doc, ROOT / "reports/results/real_data/real_data_srcs_heatmap.png", "Figure 9 Real-data SRCS heatmap", 5.0)
            continue
        if line.startswith("- "):
            flush(); p = doc.add_paragraph(f"•  {line[2:]}")
            p.paragraph_format.left_indent = Inches(0.25); p.paragraph_format.first_line_indent = Inches(-0.18)
            continue
        if re.match(r"^\d+\.\s", line):
            flush(); doc.add_paragraph(line.replace("*", "")); continue
        if line.startswith("> "):
            flush(); p = doc.add_paragraph(line[2:]); p.runs[0].italic = True; continue
        if not line.strip():
            flush()
        else:
            paragraph_buffer.append(line)
    flush()
    doc.core_properties.title = "Inferring Hidden Economic Connections from Financial Shock Propagation"
    doc.core_properties.author = "Samearth17"
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
